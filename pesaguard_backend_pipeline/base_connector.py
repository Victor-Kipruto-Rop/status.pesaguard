"""
Enterprise-grade base interface and connectors for internal ledger/order systems.

Supports multi-tenant Postgres databases, REST API ledgers, and Google Sheets integrations
with automatic retry backoff, SQL identifier sanitization, and normalized schema parsing.
"""

from __future__ import annotations

import json
import logging
import os
import re
from decimal import Decimal
from abc import ABC, abstractmethod
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, Iterable, List, Optional

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from sqlalchemy import Numeric, bindparam, create_engine, text
from normalization import (
    DEFAULT_CURRENCY,
    MAX_TRANSACTION_AMOUNT,
    NormalizationError,
    normalize_amount,
    normalize_currency,
    normalize_identifier,
    parse_decimal,
)

logger = logging.getLogger("pesaguard.connectors")

# Defense-in-depth SQL identifier allowlist validation
_SAFE_IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def _validate_identifier(name: str, kind: str) -> str:
    """Validate database object identifiers to prevent SQL injection during interpolation."""
    if not name or not _SAFE_IDENTIFIER.match(name):
        raise ValueError(
            f"Unsafe {kind} identifier rejected: {name!r}. Must contain only ASCII letters, "
            f"digits, and underscores, and must not start with a digit."
        )
    return name


class BaseConnector(ABC):
    """
    Abstract Base Class for all internal order/ledger connectors.
    
    All connectors must yield records matching the normalized PesaGuard ledger shape:
    {
        "internal_ref": str,        # Order/Invoice ID in the customer's ledger system
        "amount": Decimal | str,    # Expected monetary amount, preserved without binary-float conversion
        "phone_number": str,        # Payer phone number
        "timestamp": str,           # ISO8601 UTC timestamp
        "status": str,              # Status (e.g. "pending", "paid", "failed")
    }
    """

    @abstractmethod
    def fetch_recent_records(self, since_minutes: int = 15) -> Iterable[Dict[str, Any]]:
        """Return internal records created or modified within the last N minutes."""
        raise NotImplementedError

    def fetch_candidate_records(
        self,
        since_minutes: int = 15,
        amount: Optional[Decimal] = None,
        phone_number: Optional[str] = None,
        tolerance: Optional[Decimal] = None,
        limit: int = 500,
        reference: Optional[str] = None,
        currency: str = DEFAULT_CURRENCY,
    ) -> Iterable[Dict[str, Any]]:
        """Return all reference and amount/phone candidates without silent truncation."""
        limit = max(1, min(int(limit), 5000))
        normalized_currency = normalize_currency(currency)
        normalized_reference = (
            normalize_identifier(reference) if reference and str(reference).strip() else None
        )
        normalized_amount = (
            Decimal(normalize_amount(amount, currency=normalized_currency))
            if amount is not None
            else None
        )
        normalized_tolerance = parse_decimal(tolerance) if tolerance is not None else None
        if normalized_tolerance is not None and (
            not normalized_tolerance.is_finite()
            or normalized_tolerance < 0
            or normalized_tolerance > MAX_TRANSACTION_AMOUNT
        ):
            raise ValueError("candidate tolerance must be finite and non-negative")

        candidates = []
        for record in self.fetch_recent_records(since_minutes=since_minutes):
            if not isinstance(record, dict):
                raise ValueError("ledger connector records must be objects")
            record_reference = record.get("reference")
            reference_match = bool(
                normalized_reference
                and record_reference
                and str(record_reference).strip()
                and normalize_identifier(record_reference) == normalized_reference
            )
            amount_match = False
            if normalized_amount is not None and normalized_tolerance is not None:
                raw_record_amount = record.get("amount", record.get("TransAmount"))
                try:
                    record_currency = normalize_currency(
                        record.get("currency", record.get("Currency", DEFAULT_CURRENCY))
                    )
                    record_amount = (
                        Decimal(normalize_amount(raw_record_amount, currency=record_currency))
                        if record_currency == normalized_currency
                        else None
                    )
                except NormalizationError:
                    record_amount = None
                if record_amount is not None:
                    amount_match = (
                        abs(normalized_amount - record_amount) <= normalized_tolerance
                        and (
                            not phone_number
                            or str(record.get("phone_number") or "").strip() == phone_number
                        )
                    )
            phone_match = bool(
                phone_number
                and normalized_amount is None
                and str(record.get("phone_number") or "").strip() == phone_number
            )
            if (
                normalized_reference is None
                and normalized_amount is None
                and not phone_number
            ) or reference_match or amount_match or phone_match:
                candidates.append(record)

        if len(candidates) > limit:
            raise RuntimeError("ledger candidate set exceeds the configured limit")
        return candidates


# Default field mapping for internal ledger columns
_DEFAULT_MAPPING: Dict[str, str] = {
    "internal_ref": "internal_ref",
    "amount": "amount",
    "phone_number": "phone_number",
    "timestamp": "created_at",
    "status": "status",
}


def _merge_mapping(defaults: Dict[str, str], override: Optional[Dict[str, str]]) -> Dict[str, str]:
    """Merge custom column mappings safely over default field keys."""
    merged = dict(defaults)
    if override:
        merged.update(override)
    return merged


def _optional_ledger_fields(record: Dict[str, Any], mapping: Dict[str, str]) -> Dict[str, Any]:
    """Preserve optional matching reference and currency fields from ledger adapters."""
    fields = {}
    for name, aliases in (
        ("reference", ("reference", "merchant_ref", "payment_reference", "bill_ref_number")),
        ("currency", ("currency", "Currency")),
    ):
        mapped_name = mapping.get(name)
        value = record.get(mapped_name) if mapped_name else None
        if value in (None, ""):
            value = next((record.get(alias) for alias in aliases if record.get(alias) not in (None, "")), None)
        if value not in (None, ""):
            fields[name] = value
    return fields


class PostgresConnector(BaseConnector):
    """Connector for pilot customers whose internal orders live in PostgreSQL."""

    def __init__(self, connection_string: str, table_name: str = "orders", mapping: Optional[Dict[str, str]] = None):
        self.connection_string = connection_string
        self.table_name = table_name
        self.mapping = _merge_mapping(_DEFAULT_MAPPING, mapping)
        self._engine = None

    def _get_engine(self):
        """Lazy-initialize a single connection pool per connector instance."""
        if self._engine is None:
            if not self.connection_string:
                raise ValueError("PostgresConnector requires a valid database connection string.")
            self._engine = create_engine(
                self.connection_string,
                pool_pre_ping=True,
                pool_size=int(os.getenv("CONNECTOR_DB_POOL_SIZE", "5")),
                max_overflow=int(os.getenv("CONNECTOR_DB_MAX_OVERFLOW", "10")),
            )
        return self._engine

    def fetch_recent_records(self, since_minutes: int = 15) -> List[Dict[str, Any]]:
        """Query recent internal transactions from Postgres safely."""
        if not self.connection_string:
            raise ValueError("PostgresConnector requires a configured connection string")
        safe_table = _validate_identifier(self.table_name, "table")
        safe_columns = {
            key: _validate_identifier(value, f"column mapping for '{key}'")
            for key, value in self.mapping.items()
        }
        selected_fields = ["internal_ref", "amount", "phone_number", "timestamp", "status"]
        selected_fields.extend(
            field for field in ("reference", "currency") if self.mapping.get(field)
        )
        query = text(
            f"SELECT {', '.join(safe_columns[field] for field in selected_fields)} "
            f"FROM {safe_table} WHERE {safe_columns['timestamp']} >= :since "
            f"ORDER BY {safe_columns['timestamp']} DESC, {safe_columns['internal_ref']} ASC"
        )
        since = datetime.now(timezone.utc) - timedelta(minutes=max(0, since_minutes))
        with self._get_engine().connect() as connection:
            rows = connection.execute(query, {"since": since}).fetchall()
        records = []
        for row in rows:
            values = dict(zip(selected_fields, row))
            timestamp = values["timestamp"]
            record = {
                "internal_ref": str(values["internal_ref"] or ""),
                "amount": values["amount"],
                "phone_number": str(values["phone_number"] or ""),
                "timestamp": timestamp.isoformat() if hasattr(timestamp, "isoformat") else str(timestamp),
                "status": str(values["status"] or "pending"),
            }
            for field in ("reference", "currency"):
                if values.get(field) not in (None, ""):
                    record[field] = values[field]
            records.append(record)
        return records

    def fetch_candidate_records(
        self,
        since_minutes: int = 15,
        amount: Optional[Decimal] = None,
        phone_number: Optional[str] = None,
        tolerance: Optional[Decimal] = None,
        limit: int = 500,
        reference: Optional[str] = None,
        currency: str = DEFAULT_CURRENCY,
    ) -> List[Dict[str, Any]]:
        """Fetch amount candidates and every available row sharing a payment reference."""
        if not self.connection_string:
            raise ValueError("PostgresConnector requires a configured connection string")
        currency = normalize_currency(currency)
        safe_table = _validate_identifier(self.table_name, "table")
        safe_columns = {
            key: _validate_identifier(value, f"column mapping for '{key}'")
            for key, value in self.mapping.items()
        }
        selected_fields = ["internal_ref", "amount", "phone_number", "timestamp", "status"]
        selected_fields.extend(
            field for field in ("reference", "currency") if self.mapping.get(field)
        )
        limit = max(1, min(limit, 5000))
        params: Dict[str, Any] = {
            "since": datetime.now(timezone.utc) - timedelta(minutes=max(0, since_minutes)),
            "limit": limit + 1,
        }
        filters = []
        amount_filter = None
        if amount is not None and tolerance is not None:
            normalized_amount = Decimal(normalize_amount(amount, currency=currency))
            normalized_tolerance = parse_decimal(tolerance)
            if (
                not normalized_tolerance.is_finite()
                or normalized_tolerance < 0
                or normalized_tolerance > MAX_TRANSACTION_AMOUNT
            ):
                raise ValueError("candidate tolerance must be finite and non-negative")
            amount_filter = (
                f"{safe_columns['amount']} BETWEEN :lower_amount AND :upper_amount"
            )
            params["lower_amount"] = max(Decimal("0"), normalized_amount - normalized_tolerance)
            params["upper_amount"] = normalized_amount + normalized_tolerance
        if phone_number:
            params["phone_number"] = phone_number
            phone_filter = f"{safe_columns['phone_number']} = :phone_number"
        else:
            phone_filter = None

        amount_currency_filter = None
        if amount_filter and self.mapping.get("currency"):
            raw_currency = f"UPPER(TRIM({safe_columns['currency']}))"
            normalized_currency = (
                f"CASE {raw_currency} "
                "WHEN 'KSH' THEN 'KES' "
                "WHEN 'KSHS' THEN 'KES' "
                "WHEN 'KENYAN SHILLING' THEN 'KES' "
                "WHEN 'US DOLLAR' THEN 'USD' "
                "WHEN '$' THEN 'USD' "
                f"ELSE {raw_currency} END"
            )
            amount_currency_filter = (
                f"COALESCE(NULLIF({normalized_currency}, ''), :default_currency) = :currency"
            )
            params["currency"] = currency
            params["default_currency"] = DEFAULT_CURRENCY
        elif amount_filter and currency != DEFAULT_CURRENCY:
            amount_currency_filter = "1 = 0"

        if reference and self.mapping.get("reference"):
            dialect = self._get_engine().dialect.name
            reference_column = safe_columns["reference"]
            if dialect == "postgresql":
                normalized_reference_column = (
                    f"UPPER(REGEXP_REPLACE({reference_column}, '[[:space:]]+', '', 'g'))"
                )
            elif dialect == "sqlite":
                normalized_reference_column = f"REPLACE({reference_column}, ' ', '')"
                for codepoint in range(9, 14):
                    normalized_reference_column = (
                        f"REPLACE({normalized_reference_column}, CHAR({codepoint}), '')"
                    )
                normalized_reference_column = f"UPPER({normalized_reference_column})"
            else:
                raise RuntimeError(
                    f"reference candidate normalization is unsupported for database dialect {dialect!r}"
                )
            amount_candidates = [amount_filter] if amount_filter else []
            if phone_filter and amount_filter:
                amount_candidates.append(phone_filter)
            if amount_currency_filter:
                amount_candidates.append(amount_currency_filter)
            reference_terms = [f"{normalized_reference_column} = :reference"]
            if amount_candidates:
                reference_terms.append("(" + " AND ".join(amount_candidates) + ")")
            filters.append("(" + " OR ".join(reference_terms) + ")")
            params["reference"] = normalize_identifier(reference)
        else:
            if amount_filter:
                filters.append(amount_filter)
            if amount_currency_filter:
                filters.append(amount_currency_filter)
            if phone_filter:
                filters.append(phone_filter)
        query = text(
            f"SELECT {', '.join(safe_columns[field] for field in selected_fields)} "
            f"FROM {safe_table} WHERE {safe_columns['timestamp']} >= :since"
            + (f" AND {' AND '.join(filters)}" if filters else "")
            + f" ORDER BY {safe_columns['timestamp']} DESC, {safe_columns['internal_ref']} ASC LIMIT :limit"
        )
        if amount is not None and tolerance is not None:
            query = query.bindparams(
                bindparam("lower_amount", type_=Numeric()),
                bindparam("upper_amount", type_=Numeric()),
            )
        with self._get_engine().connect() as connection:
            rows = connection.execute(query, params).fetchall()
        if len(rows) > limit:
            raise RuntimeError("ledger candidate set exceeds the configured limit")
        records = []
        for row in rows:
            values = dict(zip(selected_fields, row))
            timestamp = values["timestamp"]
            record = {
                "internal_ref": str(values["internal_ref"] or ""),
                "amount": values["amount"],
                "phone_number": str(values["phone_number"] or ""),
                "timestamp": timestamp.isoformat() if hasattr(timestamp, "isoformat") else str(timestamp),
                "status": str(values["status"] or "pending"),
            }
            for field in ("reference", "currency"):
                if values.get(field) not in (None, ""):
                    record[field] = values[field]
            records.append(record)
        return records


class GoogleSheetsConnector(BaseConnector):
    """Production-ready Google Sheets ledger connector using Service Account Auth."""

    def __init__(
        self,
        sheet_id: str,
        worksheet_name: str = "Orders",
        credentials_json: Optional[str] = None,
        mapping: Optional[Dict[str, str]] = None,
    ):
        self.sheet_id = sheet_id
        self.worksheet_name = worksheet_name
        self.credentials_json = credentials_json or os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "")
        self.mapping = _merge_mapping(_DEFAULT_MAPPING, mapping)
        self._client = None

    def _get_client(self):
        """Initialize gspread client with Service Account credentials."""
        if self._client is not None:
            return self._client

        if not self.credentials_json:
            raise ValueError("GoogleSheetsConnector requires Service Account JSON credentials.")

        try:
            import gspread
            from google.oauth2.service_account import Credentials

            scopes = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
            if self.credentials_json.startswith("{"):
                cred_dict = json.loads(self.credentials_json)
                creds = Credentials.from_service_account_info(cred_dict, scopes=scopes)
            else:
                creds = Credentials.from_service_account_file(self.credentials_json, scopes=scopes)

            self._client = gspread.authorize(creds)
            return self._client
        except ImportError:
            logger.error("gspread or google-auth package is not installed.")
            raise RuntimeError("Missing required dependencies: 'gspread' and 'google-auth'.")

    def fetch_recent_records(self, since_minutes: int = 15) -> List[Dict[str, Any]]:
        """Fetch and parse records from the designated Google Sheet worksheet."""
        if not self.sheet_id:
            raise ValueError("GoogleSheetsConnector requires a configured sheet_id")

        try:
            client = self._get_client()
            sheet = client.open_by_key(self.sheet_id)
            worksheet = sheet.worksheet(self.worksheet_name)
            rows = worksheet.get_all_records()

            records = []
            now_ts = datetime.now(timezone.utc).timestamp()
            cutoff_ts = now_ts - (since_minutes * 60)

            amount_col = self.mapping["amount"]
            phone_col = self.mapping["phone_number"]
            time_col = self.mapping["timestamp"]
            status_col = self.mapping["status"]

            for row in rows:
                raw_time = str(row.get(time_col, "")).strip()
                record_dt = None

                # Parse timestamp formats safely
                if raw_time:
                    try:
                        record_dt = datetime.fromisoformat(raw_time.replace("Z", "+00:00"))
                        if record_dt.tzinfo is None:
                            record_dt = record_dt.replace(tzinfo=timezone.utc)
                        else:
                            record_dt = record_dt.astimezone(timezone.utc)
                    except ValueError:
                        pass

                if record_dt and record_dt.timestamp() < cutoff_ts:
                    continue

                record = {
                    "internal_ref": str(row.get(self.mapping["internal_ref"], "")),
                    "amount": row.get(amount_col, 0) or 0,
                    "phone_number": str(row.get(phone_col, "")),
                    "timestamp": record_dt.isoformat() if record_dt else raw_time,
                    "status": str(row.get(status_col, "pending")),
                }
                record.update(_optional_ledger_fields(row, self.mapping))
                records.append(record)

            logger.info("Fetched %d internal records from Google Sheet '%s'", len(records), self.sheet_id)
            return records

        except Exception as exc:
            logger.exception("Google Sheets ledger fetch failed")
            raise RuntimeError("Google Sheets ledger fetch failed") from exc


class RestConnector(BaseConnector):
    """Robust REST API ledger connector with exponential retries and error handling."""

    def __init__(
        self,
        endpoint: str,
        auth_type: str = "api_key",
        auth_value: str = "",
        mapping: Optional[Dict[str, str]] = None,
        timeout_seconds: int = 10,
    ):
        self.endpoint = endpoint
        self.auth_type = auth_type
        self.auth_value = auth_value
        self.timeout_seconds = timeout_seconds
        self.mapping = _merge_mapping(
            {"internal_ref": "id", "amount": "amount", "phone_number": "phone", "timestamp": "created_at", "status": "status"},
            mapping,
        )
        self._session = self._build_http_session()

    def _build_http_session(self) -> requests.Session:
        """Construct an HTTP session with automated retry logic on transient errors."""
        session = requests.Session()
        retries = Retry(
            total=3,
            backoff_factor=1,
            status_forcelist=[429, 500, 502, 503, 504],
            allowed_methods=["GET"],
        )
        adapter = HTTPAdapter(max_retries=retries)
        session.mount("https://", adapter)
        session.mount("http://", adapter)
        return session

    def fetch_recent_records(self, since_minutes: int = 15) -> List[Dict[str, Any]]:
        """Execute REST query to retrieve recent order payload."""
        if not self.endpoint:
            raise ValueError("RestConnector requires a configured endpoint")

        headers = {"Accept": "application/json"}
        if self.auth_type == "bearer" and self.auth_value:
            headers["Authorization"] = f"Bearer {self.auth_value}"
        elif self.auth_type == "api_key" and self.auth_value:
            headers["X-API-Key"] = self.auth_value

        params = {"since_minutes": since_minutes}

        try:
            response = self._session.get(
                self.endpoint,
                headers=headers,
                params=params,
                timeout=self.timeout_seconds,
            )
            response.raise_for_status()

            payload = response.json()
            items = payload.get("items", payload) if isinstance(payload, dict) else payload

            if not isinstance(items, list):
                raise ValueError("REST connector expected a list of ledger records")

            records = []
            for item in items:
                if not isinstance(item, dict):
                    raise ValueError("REST connector ledger rows must be objects")
                records.append({
                    "internal_ref": str(item.get(self.mapping["internal_ref"], "")),
                    "amount": item.get(self.mapping["amount"], 0) or 0,
                    "phone_number": str(item.get(self.mapping["phone_number"], "")),
                    "timestamp": str(item.get(self.mapping["timestamp"], "")),
                    "status": str(item.get(self.mapping["status"], "pending")),
                    **_optional_ledger_fields(item, self.mapping),
                })

            logger.info("Fetched %d internal records via REST connector from %s", len(records), self.endpoint)
            return records

        except requests.RequestException as exc:
            logger.error("REST ledger HTTP request failed: %s", type(exc).__name__)
            raise RuntimeError("REST ledger fetch failed") from exc
        except Exception as exc:
            logger.exception("REST ledger response processing failed")
            raise RuntimeError("REST ledger response processing failed") from exc


class ConnectorRegistry:
    """Multi-tenant connector registry supporting dynamic per-tenant configuration."""

    def __init__(self, connectors: Optional[Dict[str, BaseConnector]] = None):
        self.connectors: Dict[str, BaseConnector] = connectors or {}

    @classmethod
    def from_env(cls) -> "ConnectorRegistry":
        """Build connector registry initialized from active environment configuration."""
        connector_type = os.getenv("CONNECTOR_TYPE", "postgres").lower()
        tenant_id = os.getenv("TENANT_ID", "default")

        connector: Optional[BaseConnector] = None

        if connector_type == "postgres":
            connector = PostgresConnector(
                connection_string=os.getenv("DATABASE_URL", ""),
                table_name=os.getenv("INTERNAL_RECORDS_TABLE", "orders"),
                mapping=json.loads(os.getenv("POSTGRES_MAPPING", "{}")) if os.getenv("POSTGRES_MAPPING") else None,
            )
        elif connector_type == "google_sheets":
            connector = GoogleSheetsConnector(
                sheet_id=os.getenv("GOOGLE_SHEET_ID", ""),
                worksheet_name=os.getenv("GOOGLE_SHEET_WORKSHEET", "Orders"),
                credentials_json=os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", ""),
                mapping=json.loads(os.getenv("GOOGLE_SHEETS_MAPPING", "{}"))
                if os.getenv("GOOGLE_SHEETS_MAPPING")
                else None,
            )
        elif connector_type == "rest":
            connector = RestConnector(
                endpoint=os.getenv("REST_CONNECTOR_ENDPOINT", ""),
                auth_type=os.getenv("REST_CONNECTOR_AUTH_TYPE", "api_key"),
                auth_value=os.getenv("REST_CONNECTOR_AUTH_VALUE", ""),
                mapping=json.loads(os.getenv("REST_MAPPING", "{}")) if os.getenv("REST_MAPPING") else None,
            )

        return cls({tenant_id: connector}) if connector else cls({})

    def get_connector(self, tenant_id: str) -> Optional[BaseConnector]:
        """Retrieve the configured connector instance for a tenant."""
        return self.connectors.get(tenant_id)


class CsvLedgerConnector(BaseConnector):
    """Simple CSV-backed internal ledger reader for importing purchase orders and ledger sheets."""

    def __init__(self, csv_path: str = "", *, mapping: Optional[Dict[str, str]] = None):
        self.csv_path = csv_path
        self.mapping = _merge_mapping(_DEFAULT_MAPPING, mapping)

    def fetch_recent_records(self, since_minutes: int = 15) -> List[Dict[str, Any]]:
        if not self.csv_path:
            raise ValueError("CsvLedgerConnector requires a configured CSV path")
        try:
            import csv

            with open(self.csv_path, newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                records: List[Dict[str, Any]] = []
                for row in reader:
                    record = {
                        "internal_ref": str(row.get(self.mapping["internal_ref"], row.get("reference", ""))),
                        "amount": row.get(self.mapping["amount"], 0) or 0,
                        "phone_number": str(row.get(self.mapping["phone_number"], "") or ""),
                        "timestamp": str(row.get(self.mapping["timestamp"], "") or ""),
                        "status": str(row.get(self.mapping["status"], "pending") or "pending"),
                    }
                    record.update(_optional_ledger_fields(row, self.mapping))
                    if record["internal_ref"] or record["amount"]:
                        records.append(record)
                return records
        except Exception as exc:
            logger.exception("CSV ledger file read failed")
            raise RuntimeError("CSV ledger fetch failed") from exc


class ExcelLedgerConnector(BaseConnector):
    """Excel-backed internal ledger connector for .xlsx and .xls import work."""

    def __init__(self, source_path: str = "", *, mapping: Optional[Dict[str, str]] = None):
        self.source_path = source_path
        self.mapping = _merge_mapping(_DEFAULT_MAPPING, mapping)

    def fetch_recent_records(self, since_minutes: int = 15) -> List[Dict[str, Any]]:
        if not self.source_path:
            raise ValueError("ExcelLedgerConnector requires a configured source path")
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(self.source_path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                if not headers:
                    return []
                normalized = [str(value).strip() for value in headers]
                records: List[Dict[str, Any]] = []
                for values in rows:
                    row = {header: value for header, value in zip(normalized, values) if header}
                    if not row:
                        continue
                    record = {
                        "internal_ref": str(row.get(self.mapping["internal_ref"], row.get("reference", ""))),
                        "amount": row.get(self.mapping["amount"], 0) or 0,
                        "phone_number": str(row.get(self.mapping["phone_number"], "") or ""),
                        "timestamp": str(row.get(self.mapping["timestamp"], "") or ""),
                        "status": str(row.get(self.mapping["status"], "pending") or "pending"),
                    }
                    record.update(_optional_ledger_fields(row, self.mapping))
                    if record["internal_ref"] or record["amount"]:
                        records.append(record)
                return records
            finally:
                workbook.close()
        except Exception as exc:
            logger.exception("Excel ledger file read failed")
            raise RuntimeError("Excel ledger fetch failed") from exc
