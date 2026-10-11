from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping, Sequence


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _strip_non_alnum(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ""
    return re.sub(r"[^A-Za-z0-9]+", "", text).upper()


def _normalize_reference(value: Any) -> str:
    return _strip_non_alnum(value)


def _normalize_phone(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ""
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits:
        return ""
    if digits.startswith("254"):
        return "+" + digits
    if digits.startswith("0"):
        return "+254" + digits[1:]
    if digits.startswith("7"):
        return "+254" + digits
    if digits.startswith("+254"):
        return "+" + digits[1:]
    return "+" + digits


def _coerce_decimal(value: Any) -> Decimal:
    if value is None or value == "":
        raise ValueError("amount is missing")
    text = str(value).strip().replace(",", "")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"invalid numeric amount: {value!r}") from exc


def _coerce_timestamp(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        value = value.astimezone(timezone.utc) if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return value
    text = str(value).strip()
    if not text:
        return None
    candidates = [text]
    if text.endswith("Z"):
        candidates.append(text[:-1] + "+00:00")
    for candidate in candidates:
        try:
            parsed = datetime.fromisoformat(candidate)
            if parsed.tzinfo is None:
                return parsed.replace(tzinfo=timezone.utc)
            return parsed.astimezone(timezone.utc)
        except ValueError:
            pass
    if len(text) == 14 and text.isdigit():
        try:
            return datetime.strptime(text, "%Y%m%d%H%M%S").replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    return None


@dataclass
class PurchaseOrder:
    po_number: str
    vendor: str = ""
    amount: Decimal | None = None
    currency: str = "KES"
    status: str = "pending"
    phone_number: str = ""
    timestamp: datetime | None = None
    reference: str | None = None
    raw: dict[str, Any] = field(default_factory=dict)

    @classmethod
    def from_mapping(cls, values: Mapping[str, Any]) -> "PurchaseOrder":
        if not isinstance(values, Mapping):
            raise ValueError("purchase order rows must be mappings")
        row = {str(k): v for k, v in values.items()}
        po_number = _clean_text(
            row.get("po_number")
            or row.get("purchase_order")
            or row.get("purchase_order_number")
            or row.get("po")
            or row.get("order_number")
            or row.get("orderid")
            or row.get("reference")
        )
        if not po_number:
            raise ValueError("purchase order number is required")
        amount_value = row.get("amount")
        amount = _coerce_decimal(amount_value) if amount_value not in (None, "") else None
        return cls(
            po_number=po_number,
            vendor=_clean_text(row.get("vendor") or row.get("supplier") or row.get("merchant") or row.get("company")),
            amount=amount,
            currency=_clean_text(row.get("currency") or "KES").upper() or "KES",
            status=_clean_text(row.get("status") or "pending").lower(),
            phone_number=_normalize_phone(row.get("phone_number") or row.get("phone") or row.get("msisdn") or row.get("customer_phone")),
            timestamp=_coerce_timestamp(row.get("timestamp") or row.get("date") or row.get("created_at") or row.get("trans_time")),
            reference=_normalize_reference(row.get("reference") or row.get("ref") or row.get("po_number") or row.get("po")),
            raw=dict(row),
        )


@dataclass
class LedgerEntry:
    entry_id: str
    amount: Decimal | None = None
    phone_number: str = ""
    timestamp: datetime | None = None
    status: str = "pending"
    reference: str | None = None
    source: str = "manual"
    raw: dict[str, Any] = field(default_factory=dict)

    @classmethod
    def from_mapping(cls, values: Mapping[str, Any], *, source: str = "manual") -> "LedgerEntry":
        if not isinstance(values, Mapping):
            raise ValueError("ledger rows must be mappings")
        row = {str(k): v for k, v in values.items()}
        entry_id = _clean_text(
            row.get("entry_id")
            or row.get("id")
            or row.get("transaction_id")
            or row.get("reference")
            or row.get("po_number")
            or row.get("ledger_id")
            or "ledger-entry"
        )
        amount_value = row.get("amount") or row.get("trans_amount") or row.get("gross_amount")
        amount = _coerce_decimal(amount_value) if amount_value not in (None, "") else None
        return cls(
            entry_id=entry_id,
            amount=amount,
            phone_number=_normalize_phone(row.get("phone_number") or row.get("phone") or row.get("msisdn") or row.get("customer_phone")),
            timestamp=_coerce_timestamp(row.get("timestamp") or row.get("date") or row.get("trans_time") or row.get("created_at") or row.get("transaction_date")),
            status=_clean_text(row.get("status") or "pending").lower(),
            reference=_normalize_reference(row.get("reference") or row.get("ref") or row.get("po_number") or row.get("transaction_ref")),
            source=source,
            raw=dict(row),
        )


def _iter_rows(rows: Iterable[Mapping[str, Any]]) -> Iterator[Mapping[str, Any]]:
    for row in rows:
        if isinstance(row, Mapping):
            yield row


def _dialect_sniff(text: str) -> str:
    sample = text.lstrip("\ufeff")
    if ";" in sample and "," not in sample:
        return ";"
    if "," in sample and ";" not in sample:
        return ","
    return ","


def parse_csv_records(content: str | bytes, *, dialect: str | None = None) -> list[dict[str, Any]]:
    if isinstance(content, bytes):
        text = content.decode("utf-8-sig")
    else:
        text = content
    sample = text.lstrip("\ufeff")
    if not sample.strip():
        raise ValueError("CSV import is empty")
    delimiter = dialect or _dialect_sniff(sample)
    reader = csv.DictReader(io.StringIO(sample), delimiter=delimiter)
    if not reader.fieldnames:
        raise ValueError("CSV import must include a header row")
    rows = []
    for row in reader:
        if row and any((value is not None and str(value).strip()) for value in row.values()):
            rows.append({key: (value if value is not None else "") for key, value in row.items()})
    return rows


def parse_excel_records(content: bytes | str, *, sheet_name: str | None = None) -> list[dict[str, Any]]:
    if isinstance(content, str):
        content = content.encode("utf-8")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for Excel imports") from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if sheet_name:
            selection = workbook[sheet_name]
        else:
            selection = workbook.active
        rows = selection.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("Excel sheet is empty")
        normalized_headers = [str(value).strip() for value in headers]
        records: list[dict[str, Any]] = []
        for values in rows:
            item = {header: value for header, value in zip(normalized_headers, values) if header}
            if item and any((value is not None and str(value).strip()) for value in item.values()):
                records.append(item)
        return records
    finally:
        workbook.close()


def import_records_from_file(path: str | Path, *, mime_type: str | None = None) -> list[dict[str, Any]]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if mime_type == "text/csv" or suffix in {".csv"}:
        return parse_csv_records(file_path.read_bytes())
    if mime_type in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel"} or suffix in {".xlsx", ".xls"}:
        return parse_excel_records(file_path.read_bytes())
    raise ValueError(f"unsupported file type for import: {suffix or mime_type}")


def to_purchase_orders(rows: Sequence[Mapping[str, Any]]) -> list[PurchaseOrder]:
    orders: list[PurchaseOrder] = []
    for row in _iter_rows(rows):
        try:
            orders.append(PurchaseOrder.from_mapping(row))
        except ValueError:
            continue
    return orders


def to_ledger_entries(rows: Sequence[Mapping[str, Any]], *, source: str = "manual") -> list[LedgerEntry]:
    result: list[LedgerEntry] = []
    for row in _iter_rows(rows):
        try:
            result.append(LedgerEntry.from_mapping(row, source=source))
        except ValueError:
            continue
    return result


__all__ = [
    "PurchaseOrder",
    "LedgerEntry",
    "parse_csv_records",
    "parse_excel_records",
    "import_records_from_file",
    "to_purchase_orders",
    "to_ledger_entries",
]
